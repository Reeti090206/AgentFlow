import ast
import operator as op
import os
import pandas as pd
from pypdf import PdfReader
from ollama_client import query_ollama
import memory_db

# Safe calculator implementation using AST parsing
SUPPORTED_OPERATORS = {
    ast.Add: op.add,
    ast.Sub: op.sub,
    ast.Mult: op.mul,
    ast.Div: op.truediv,
    ast.Pow: op.pow,
    ast.USub: op.neg,
    ast.UAdd: lambda x: x
}

def _eval_ast_node(node):
    if isinstance(node, ast.Num):  # Compatibility for older Python
        return node.n
    elif isinstance(node, ast.Constant):  # Python 3.8+
        return node.value
    elif isinstance(node, ast.BinOp):
        op_type = type(node.op)
        if op_type in SUPPORTED_OPERATORS:
            return SUPPORTED_OPERATORS[op_type](_eval_ast_node(node.left), _eval_ast_node(node.right))
        raise ValueError(f"Operator {op_type.__name__} not supported.")
    elif isinstance(node, ast.UnaryOp):
        op_type = type(node.op)
        if op_type in SUPPORTED_OPERATORS:
            return SUPPORTED_OPERATORS[op_type](_eval_ast_node(node.operand))
        raise ValueError(f"Operator {op_type.__name__} not supported.")
    raise TypeError(f"Unsupported expression component: {type(node).__name__}")

def calculator(expression):
    """Safely evaluates a mathematical expression string.
    
    Example: calculator("2 * (3 + 4) / 5")
    """
    clean_expr = str(expression).replace(" ", "")
    try:
        parsed = ast.parse(clean_expr, mode='eval')
        result = _eval_ast_node(parsed.body)
        return f"Calculation Result: {result}"
    except Exception as e:
        return f"Calculator Error: Could not evaluate '{expression}'. Details: {str(e)}"

def read_file(filepath):
    """Extracts text content from local TXT or PDF files."""
    if not os.path.exists(filepath):
        return f"File Error: File not found at '{filepath}'."
    
    ext = os.path.splitext(filepath)[1].lower()
    
    if ext == ".pdf":
        try:
            reader = PdfReader(filepath)
            text = ""
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text += f"--- Page {idx+1} ---\n{page_text}\n"
            if not text.strip():
                return "File Content: (PDF file appears to be empty or image-only)"
            return text
        except Exception as e:
            return f"File Error: Could not parse PDF file. Details: {str(e)}"
            
    elif ext == ".docx":
        try:
            import docx
            doc = docx.Document(filepath)
            text = ""
            for para in doc.paragraphs:
                if para.text:
                    text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text]
                    if row_text:
                        text += " | ".join(row_text) + "\n"
            if not text.strip():
                return "File Content: (DOCX file appears to be empty)"
            return text
        except Exception as e:
            return f"File Error: Could not parse DOCX file. Details: {str(e)}"
            
    elif ext == ".csv":
        return f"File notice: '{filepath}' is a CSV database. Please use the analyze_csv tool to examine it."
        
    else:
        # Treat as plain text
        try:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
            return content
        except Exception as e:
            return f"File Error: Could not read file. Details: {str(e)}"

def summarize_text(text, model=None):
    """Generates a concise summary of the provided text using local Ollama LLM."""
    if not text or not text.strip():
        return "Text Summary: Nothing to summarize."
        
    # Truncate text if it's extremely long to avoid context window overflow in tiny local LLMs
    max_char_len = 12000
    truncated = text[:max_char_len] + ("... (truncated for summary)" if len(text) > max_char_len else "")
    
    prompt = f"Provide a brief, clear, and structured summary of the following text in 3-5 key bullet points:\n\n{truncated}"
    system_prompt = "You are a professional text summarizer. Provide only the concise bullet-point summary, with no chatty intro or outro."
    
    success, result = query_ollama(prompt, system_prompt=system_prompt, model=model, temperature=0.3)
    if success:
        return f"Text Summary:\n{result}"
    else:
        return f"Summarizer Error: {result}"

def analyze_csv(filepath):
    """Loads a CSV file, computes shape, column list, missing values, descriptive statistics, and a data preview."""
    if not os.path.exists(filepath):
        return f"CSV Error: File not found at '{filepath}'."
        
    try:
        # Load the CSV
        df = pd.read_csv(filepath)
        
        # Get shape
        rows, cols = df.shape
        
        # Get column details and types
        col_types = df.dtypes.to_dict()
        col_list_str = ", ".join([f"{name} ({t})" for name, t in col_types.items()])
        
        # Missing values
        missing_vals = df.isnull().sum().to_dict()
        missing_str = ", ".join([f"{k}: {v}" for k, v in missing_vals.items() if v > 0])
        if not missing_str:
            missing_str = "None"
            
        # Describe numeric columns
        numeric_desc = ""
        numeric_cols = df.select_dtypes(include=['number']).columns.tolist()
        if numeric_cols:
            desc = df[numeric_cols].describe()
            numeric_desc = "\nNumeric Columns Statistics:\n" + desc.to_string()
            
        # Preview rows
        preview = df.head(3).to_string()
        
        result_str = (
            f"CSV Analysis Report for '{os.path.basename(filepath)}':\n"
            f"- Total Rows: {rows}\n"
            f"- Total Columns: {cols}\n"
            f"- Columns: {col_list_str}\n"
            f"- Missing Values: {missing_str}\n"
            f"{numeric_desc}\n\n"
            f"First 3 Rows Preview:\n{preview}"
        )
        return result_str
        
    except Exception as e:
        return f"CSV Error: Could not analyze CSV dataset. Details: {str(e)}"

def store_memory(username, key, value):
    """Directly invokes the memory database to store/update user memory."""
    success, msg = memory_db.store_memory(username, key, value)
    return msg

def get_memory(username, key):
    """Directly invokes the memory database to get user memory."""
    val = memory_db.get_memory(username, key)
    if val:
        return f"Memory Retrieved: '{key}' = '{val}'"
    return f"Memory Notice: Key '{key}' not found in memories."
